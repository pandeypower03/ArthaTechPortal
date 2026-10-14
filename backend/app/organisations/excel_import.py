from decimal import Decimal

from django.contrib.auth import get_user_model
from django.db import transaction
from openpyxl import load_workbook

from .models import Organisation, Plan, Project, Team, TeamMember

User = get_user_model()
FEATURES = ["whatsapp_alerts", "bill_ocr", "tally_export", "multi_project_dashboard"]


def text(row, key):
    value = row.get(key)
    if isinstance(value, float) and value.is_integer():   # Excel gives 9800000001.0
        value = int(value)
    return str(value if value is not None else "").strip()


def paise(row, key):
    return int((Decimal(text(row, key) or "0") * 100).to_integral_value())   # rupees -> paise, no float


def import_excel(file):
    """Read the 'Upload' sheet and save it. All or nothing. Raises ValueError("Row 5: ...")."""
    try:
        lines = load_workbook(file, read_only=True, data_only=True)["Upload"].iter_rows(values_only=True)
        headings = [str(h).strip() for h in next(lines)]
    except Exception:
        raise ValueError("Cannot read the file. Please use the upload template.")

    counts = {"houses": 0, "projects": 0, "teams": 0, "members": 0}   # only NEW records are counted
    with transaction.atomic():
        for number, values in enumerate(lines, start=2):
            if not any(values):
                continue
            try:
                save_row(dict(zip(headings, values)), counts)
            except Exception as error:
                raise ValueError(f"Row {number}: {error}")
    return counts


def save_row(row, counts):
    plan = Plan.objects.filter(name__iexact=text(row, "plan_name")).first()
    if not text(row, "organisation_name") or not plan:
        raise ValueError("organisation_name is empty or plan_name does not exist")

    org, new = Organisation.objects.get_or_create(
        name=text(row, "organisation_name"),
        defaults={"plan": plan, **{f: text(row, f).lower() == "yes" for f in FEATURES}},
    )
    counts["houses"] += new

    if not text(row, "project_name"):
        return
    if text(row, "project_type").lower() not in Project.Type.values:
        raise ValueError("project_type must be movie, ott_series, ad or other")
    project, new = Project.objects.get_or_create(
        organisation=org, name=text(row, "project_name"),
        defaults={"project_type": text(row, "project_type").lower()},
    )
    counts["projects"] += new

    if not text(row, "team_name"):
        return
    team, new = Team.objects.get_or_create(
        project=project, name=text(row, "team_name"),
        defaults={"budget_paise": paise(row, "team_budget_rupees"), "advance_paise": paise(row, "advance_rupees")},
    )
    counts["teams"] += new

    if not text(row, "member_name"):
        return
    user = None
    if text(row, "login_email"):
        user = User.objects.filter(email=text(row, "login_email").lower(), is_superuser=False).first()
        if not user or user.organisation_id not in (None, org.id):
            raise ValueError("login_email has no login, or it belongs to another house")
        user.organisation = org
        user.save(update_fields=["organisation"])
    member, new = TeamMember.objects.get_or_create(
        team=team, name=text(row, "member_name"), phone=text(row, "phone"),
        defaults={"job_title": text(row, "job_title"), "is_team_head": text(row, "is_team_head").lower() == "yes", "user": user},
    )
    counts["members"] += new